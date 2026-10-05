"""Извлечение текста из документов (PDF, DOCX, TXT, MD...)."""
from __future__ import annotations

import io
import logging

log = logging.getLogger("tldr.extract")

TEXT_EXTENSIONS = {"txt", "md", "markdown", "csv", "json", "html", "htm"}
DOC_EXTENSIONS = {"pdf", "docx"}


def extract_text(filename: str, data: bytes) -> str:
    """Вернуть текстовое содержимое файла. Бросает ValueError, если формат не поддерживается."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in TEXT_EXTENSIONS:
        return data.decode("utf-8", errors="ignore")

    if ext == "pdf":
        return _extract_pdf(data)

    if ext == "docx":
        return _extract_docx(data)

    raise ValueError(
        f"Формат «.{ext}» не поддерживается. "
        "Присылай PDF, DOCX или обычный текст (TXT/MD)."
    )


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("Для PDF нужен пакет pypdf: pip install pypdf")

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception as e:  # повреждённая страница — пропускаем
            log.warning("pdf page parse error: %s", e)
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError(
            "В PDF не нашлось текста — похоже, это скан с картинками. "
            "Пока умею только текстовые PDF."
        )
    return text


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise RuntimeError("Для DOCX нужен пакет python-docx: pip install python-docx")

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    # таблицы
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    text = "\n".join(lines).strip()
    if not text:
        raise ValueError("В документе не нашлось текста.")
    return text
